"""Real DOCX fixtures containing native Office Math (OMML) structures."""

from __future__ import annotations

import io
from typing import Any

from docx import Document as DocxDocument
from docx.oxml import OxmlElement


def _run(text: str) -> Any:
    run = OxmlElement("m:r")
    leaf = OxmlElement("m:t")
    leaf.text = text
    run.append(leaf)
    return run


def _fraction(numerator: str, denominator: str) -> Any:
    fraction = OxmlElement("m:f")
    num = OxmlElement("m:num")
    num.append(_run(numerator))
    den = OxmlElement("m:den")
    den.append(_run(denominator))
    fraction.extend((num, den))
    return fraction


def _superscript(base: str, exponent: str) -> Any:
    superscript = OxmlElement("m:sSup")
    expression = OxmlElement("m:e")
    expression.append(_run(base))
    sup = OxmlElement("m:sup")
    sup.append(_run(exponent))
    superscript.extend((expression, sup))
    return superscript


def _subscript(base: str, subscript: str) -> Any:
    node = OxmlElement("m:sSub")
    expression = OxmlElement("m:e")
    expression.append(_run(base))
    sub = OxmlElement("m:sub")
    sub.append(_run(subscript))
    node.extend((expression, sub))
    return node


def _radical(radicand: str, *, degree: str | None = None) -> Any:
    radical = OxmlElement("m:rad")
    properties = OxmlElement("m:radPr")
    if degree is None:
        degree_hidden = OxmlElement("m:degHide")
        degree_hidden.set("{http://schemas.openxmlformats.org/officeDocument/2006/math}val", "1")
        properties.append(degree_hidden)
    radical.append(properties)
    degree_node = OxmlElement("m:deg")
    if degree is not None:
        degree_node.append(_run(degree))
    radical.append(degree_node)
    expression = OxmlElement("m:e")
    expression.append(_run(radicand))
    radical.append(expression)
    return radical


def _delimiter(*expressions: Any, begin: str = "[", end: str = "]", separator: str = "; ") -> Any:
    delimiter = OxmlElement("m:d")
    properties = OxmlElement("m:dPr")
    begin_node = OxmlElement("m:begChr")
    begin_node.set("{http://schemas.openxmlformats.org/officeDocument/2006/math}val", begin)
    end_node = OxmlElement("m:endChr")
    end_node.set("{http://schemas.openxmlformats.org/officeDocument/2006/math}val", end)
    separator_node = OxmlElement("m:sepChr")
    separator_node.set("{http://schemas.openxmlformats.org/officeDocument/2006/math}val", separator)
    properties.extend((begin_node, end_node, separator_node))
    delimiter.extend((properties, *expressions))
    return delimiter


def native_equation_docx_bytes() -> bytes:
    """Build a valid DOCX with mixed, nested native Word equations."""
    doc = DocxDocument()

    fraction_paragraph = doc.add_paragraph("Fraction: ")
    fraction_math = OxmlElement("m:oMath")
    fraction_math.extend((_fraction("1", "2"), _run(" + "), _fraction("1", "3")))
    fraction_paragraph._p.append(fraction_math)
    fraction_paragraph.add_run(".")

    power_paragraph = doc.add_paragraph("Power: ")
    power_math = OxmlElement("m:oMath")
    power_math.append(_superscript("2", "3"))
    power_paragraph._p.append(power_math)
    power_paragraph.add_run(".")

    root_paragraph = doc.add_paragraph("Root: ")
    root_math = OxmlElement("m:oMath")
    root_math.append(_radical("9"))
    root_paragraph._p.append(root_math)
    root_paragraph.add_run(".")

    scripts_paragraph = doc.add_paragraph("Scripts: ")
    scripts_math = OxmlElement("m:oMath")
    scripts_math.extend((_subscript("a", "i"), _run(" + ")))
    nested_fraction = OxmlElement("m:f")
    nested_numerator = OxmlElement("m:num")
    nested_numerator.append(_superscript("x", "2"))
    nested_denominator = OxmlElement("m:den")
    nested_denominator.append(_radical("y"))
    nested_fraction.extend((nested_numerator, nested_denominator))
    scripts_math.append(nested_fraction)
    scripts_paragraph._p.append(scripts_math)
    scripts_paragraph.add_run(".")

    indexed_root_paragraph = doc.add_paragraph("Indexed root: ")
    indexed_root_math = OxmlElement("m:oMath")
    indexed_root_math.append(_radical("8", degree="3"))
    indexed_root_paragraph._p.append(indexed_root_math)
    indexed_root_paragraph.add_run(".")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def unsupported_equation_docx_bytes() -> bytes:
    """Build a DOCX containing an intentionally unknown OMML node."""
    doc = DocxDocument()
    paragraph = doc.add_paragraph("Unsupported: ")
    math = OxmlElement("m:oMath")
    unknown = OxmlElement("m:mystery")
    unknown.append(_run("12"))
    math.append(unknown)
    paragraph._p.append(math)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def equation_chain_docx_bytes() -> bytes:
    """Build a DOCX covering inline, display, delimiter, and table math."""
    doc = DocxDocument()

    inline_paragraph = doc.add_paragraph("Inline before ")
    inline_math = OxmlElement("m:oMath")
    inline_math.append(_fraction("1", "2"))
    inline_paragraph._p.append(inline_math)
    inline_paragraph.add_run(" after.")

    display_paragraph = doc.add_paragraph("Display before")
    display_container = OxmlElement("m:oMathPara")
    display_math = OxmlElement("m:oMath")
    display_expression = OxmlElement("m:e")
    display_expression.extend((_run("x"), _run(" + "), _run("1")))
    display_expression_2 = OxmlElement("m:e")
    display_expression_2.append(_run("y"))
    display_math.append(_delimiter(display_expression, display_expression_2))
    display_container.append(display_math)
    display_paragraph._p.append(display_container)
    display_paragraph.add_run("after.")

    table = doc.add_table(rows=1, cols=1)
    cell_paragraph = table.cell(0, 0).paragraphs[0]
    cell_paragraph.add_run("Cell power: ")
    cell_math = OxmlElement("m:oMath")
    cell_math.append(_superscript("y", "2"))
    cell_paragraph._p.append(cell_math)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


EXPECTED_EQUATION_PARAGRAPHS = (
    "Fraction: 1/2 + 1/3.",
    "Power: 2^3.",
    "Root: sqrt(9).",
    "Scripts: a_i + x^2/sqrt(y).",
    "Indexed root: root(3, 8).",
)

EXPECTED_CHAIN_PARAGRAPHS = (
    "Inline before 1/2 after.",
    "Display before\n[x + 1; y]\nafter.",
    "Cell power: y^2",
)
