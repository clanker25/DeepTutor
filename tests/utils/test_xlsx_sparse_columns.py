"""Sparse XLSX cells must retain their column positions in the raw fallback."""

import io
import zipfile

from openpyxl import Workbook, load_workbook
import pytest

from deeptutor.utils.document_extractor import _extract_xlsx_ooxml


@pytest.mark.parametrize(
    "cells",
    [
        {"A1": "Name", "C1": "Score", "A2": "Alice", "B2": "pending", "C2": 98},
        {"B1": "Term", "C1": "Meaning", "A2": "1", "B2": "Mass", "C2": "matter"},
        {"AB1": "Value", "AD1": "Units", "AB2": 12, "AC2": "+/- 1", "AD2": "kg"},
        {"XFD1": "Last column", "A2": "First column"},
    ],
)
def test_sparse_cells_match_the_real_workbook_columns(cells: dict[str, object]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    for address, value in cells.items():
        sheet[address] = value
    stream = io.BytesIO()
    workbook.save(stream)
    data = stream.getvalue()
    actual_workbook = load_workbook(io.BytesIO(data), data_only=True)
    try:
        rows = [
            "\t".join("" if cell is None else str(cell) for cell in row).rstrip("\t")
            for row in actual_workbook.active.iter_rows(values_only=True)
        ]
    finally:
        actual_workbook.close()
    expected = "--- Sheet: Results ---\n" + "\n".join(rows)
    assert _extract_xlsx_ooxml(data, "results.xlsx") == expected


@pytest.mark.parametrize("reference", ["", "A0", "XFE1", "ZZZZZZZZ999999999"])
def test_invalid_cell_coordinates_keep_bounded_sequential_fallback(reference: str) -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            '<worksheet><sheetData><row><c r="' + reference + '"><v>12</v></c>'
            "<c><v>34</v></c></row></sheetData></worksheet>",
        )
    assert _extract_xlsx_ooxml(stream.getvalue(), "legacy.xlsx") == "--- Sheet: sheet1 ---\n12\t34"
