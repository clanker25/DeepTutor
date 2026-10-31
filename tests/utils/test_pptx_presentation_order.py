"""OOXML fallback order must agree with the actual presentation slide list."""

import io
import zipfile

from pptx import Presentation
import pytest

from deeptutor.utils.document_extractor import _extract_pptx_ooxml
from deeptutor.utils.document_images import _extract_pptx_rich_ooxml


def _reordered_deck(*, orphan: bool) -> bytes:
    presentation = Presentation()
    for title in ("First created", "Second created", "Third created"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
    # PowerPoint reorders slide relationships without renaming their parts.
    slide_ids = presentation.slides._sldIdLst
    last = slide_ids[-1]
    slide_ids.remove(last)
    slide_ids.insert(0, last)
    stream = io.BytesIO()
    presentation.save(stream)
    if not orphan:
        return stream.getvalue()
    result = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(result, "w") as target:
        for member in source.infolist():
            target.writestr(member, source.read(member))
        target.writestr("ppt/slides/slide9.xml", source.read("ppt/slides/slide2.xml"))
    return result.getvalue()


@pytest.mark.parametrize("orphan", [False, True])
def test_rich_fallback_uses_presentation_order_not_part_names(orphan: bool) -> None:
    data = _reordered_deck(orphan=orphan)
    expected = tuple(slide.shapes.title.text for slide in Presentation(io.BytesIO(data)).slides)
    assert expected == ("Third created", "First created", "Second created")
    assert _extract_pptx_rich_ooxml(data).slides == expected


@pytest.mark.parametrize("orphan", [False, True])
def test_text_fallback_uses_presentation_order_not_part_names(orphan: bool) -> None:
    data = _reordered_deck(orphan=orphan)
    expected = [slide.shapes.title.text for slide in Presentation(io.BytesIO(data)).slides]
    text = _extract_pptx_ooxml(data, "reordered.pptx")
    assert text == "\n\n".join(
        f"--- Slide {index} ---\n{title}" for index, title in enumerate(expected, 1)
    )


@pytest.mark.parametrize("malformed", [False, True])
def test_metadata_free_fallback_remains_available(malformed: bool) -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_reordered_deck(orphan=False))) as source:
        with zipfile.ZipFile(output, "w") as target:
            for member in source.infolist():
                if member.filename != "ppt/presentation.xml":
                    target.writestr(member, source.read(member))
                elif malformed:
                    target.writestr(member, "<broken")
    assert _extract_pptx_rich_ooxml(output.getvalue()).slides == (
        "First created",
        "Second created",
        "Third created",
    )


def test_empty_presentation_does_not_resurrect_an_orphan_slide() -> None:
    empty = io.BytesIO()
    Presentation().save(empty)
    output = io.BytesIO()
    with zipfile.ZipFile(empty) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            target.writestr(member, source.read(member))
        with zipfile.ZipFile(io.BytesIO(_reordered_deck(orphan=False))) as populated:
            target.writestr("ppt/slides/slide1.xml", populated.read("ppt/slides/slide1.xml"))
    data = output.getvalue()
    assert len(Presentation(io.BytesIO(data)).slides) == 0
    assert _extract_pptx_rich_ooxml(data).slides == ()
    assert _extract_pptx_ooxml(data, "empty.pptx") == ""


def test_external_slide_relationships_are_not_replaced_by_unreferenced_parts() -> None:
    from xml.etree import ElementTree

    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_reordered_deck(orphan=False))) as source:
        with zipfile.ZipFile(output, "w") as target:
            for member in source.infolist():
                data = source.read(member)
                if member.filename == "ppt/_rels/presentation.xml.rels":
                    relationships = ElementTree.fromstring(data)
                    for relationship in relationships:
                        if relationship.get("Type", "").endswith("/slide"):
                            relationship.set("TargetMode", "External")
                    data = ElementTree.tostring(relationships)
                target.writestr(member, data)
    assert _extract_pptx_rich_ooxml(output.getvalue()).slides == ()
    assert _extract_pptx_ooxml(output.getvalue(), "external.pptx") == ""
