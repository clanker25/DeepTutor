"""The OOXML fallback must retain picture shapes outside text paragraphs."""

import importlib
import io
import random
import sys
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from deeptutor.utils.document_images import ImageBudget, build_marker, extract_pptx_rich


def _presentation(*, pictures: int = 1, grouped: bool = False) -> bytes:
    """Create a real PowerPoint archive with text before and after each picture."""
    image = Image.frombytes("RGB", (100, 100), random.Random(31).randbytes(30000))
    png = io.BytesIO()
    image.save(png, format="PNG")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
    shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = "Before picture"
    for index in range(pictures):
        shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(2 + index), Inches(2))
    shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1)).text = "After picture"
    result = io.BytesIO()
    deck.save(result)
    return result.getvalue()


@pytest.mark.parametrize("mode", ["missing", "failure"])
@pytest.mark.parametrize("grouped", [False, True], ids=["ordinary", "grouped"])
def test_public_fallback_keeps_picture_shapes(
    monkeypatch: pytest.MonkeyPatch, mode: str, grouped: bool
) -> None:
    """Both fallback entry points preserve images and source ordering."""
    data = _presentation(grouped=grouped)
    expected = extract_pptx_rich(data)
    assert len(expected.collection.images) == 1
    if mode == "missing":
        monkeypatch.setitem(sys.modules, "pptx", None)
    else:
        module = importlib.import_module("pptx")

        def fail(_data: object) -> None:
            raise ValueError("Synthetic parser failure")

        monkeypatch.setattr(module, "Presentation", fail)

    actual = extract_pptx_rich(data)
    assert actual.collection.images == expected.collection.images
    marker = build_marker(actual.collection.images[0])
    text = actual.slides[0]
    assert text.index("Before picture") < text.index(marker) < text.index("After picture")


def test_fallback_applies_image_budget_once_per_picture(monkeypatch: pytest.MonkeyPatch) -> None:
    """Out-of-paragraph pictures must use the same accounting as the normal path."""
    data = _presentation(pictures=2)
    monkeypatch.setitem(sys.modules, "pptx", None)
    rich = extract_pptx_rich(data, budget=ImageBudget(max_images=1))
    assert len(rich.collection.images) == 1
    assert rich.collection.skipped_budget == 1
    assert rich.slides[0].count(build_marker(rich.collection.images[0])) == 1


def test_fallback_keeps_text_when_picture_target_is_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A missing picture part must not hide the surrounding slide text."""
    original = zipfile.ZipFile(io.BytesIO(_presentation()))
    output = io.BytesIO()
    with original, zipfile.ZipFile(output, "w") as archive:
        for member in original.infolist():
            if not member.filename.startswith("ppt/media/"):
                archive.writestr(member, original.read(member.filename))
    monkeypatch.setitem(sys.modules, "pptx", None)
    rich = extract_pptx_rich(output.getvalue())
    assert not rich.collection.images
    assert rich.slides == ("Before picture\nAfter picture",)
